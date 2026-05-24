from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

excel_file_name = "test prices.xlsx"
item_categories= ['xxx', 'yyy', 'zzz']
price_categories=['A','B', 'C','D','E','F','G','H','I','K']

color_map={'A':'red',
           'B':'green', 'C':'green','D':'green',
           'E':'blue', 'F':'blue','G':'blue','H':'blue','I':'blue','K':'blue'}

wb = load_workbook(excel_file_name)
ws = wb['Charts']


row_min = 2
row_max= 12
for category in item_categories:
    chart = BarChart()
    # Определение данных для диаграммы
    chart_values = Reference(worksheet=ws,
                             min_row=row_min+1,
                             max_row=row_max,
                             min_col=2,
                             max_col=2)
    chart_cats = Reference(ws, min_col=1, min_row=row_min + 1, max_row=row_max)
    # Добавление данных в диаграмму
    chart.add_data(chart_values)
    chart.set_categories(chart_cats)
    # Настройки внешнего вида
    chart.title="Percentage for " + category
    chart.legend=None
    chart.y_axis.delete = False
    chart.y_axis.scaling.max=100
    chart.x_axis.delete = False

    chart.varyColors=False
    # Размещение и габариты
    from_marker = AnchorMarker( col=4, row=row_min)
    to_marker = AnchorMarker(col=10, row=row_max)
    chart.anchor = TwoCellAnchor(_from=from_marker, to=to_marker)

    # Подписи к данным
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showSerName = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showLegendKey = False

    ws.add_chart(chart)
    row_min += 13
    row_max += 13

wb.save(excel_file_name)
